"""PDF/DOCX/text -> text units with source references. Files are never executed or kept.
Scanned PDF pages are read with OCR (Tesseract) when it is installed; see docs/INTEGRATIONS.md."""
import io
import zipfile

from app.core.config import get_settings
from app.core.errors import AppError

MAX_OCR_PAGES = 30


def _kind(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf") and data[:5] == b"%PDF-":
        return "pdf"
    if name.endswith(".docx") and data[:4] == b"PK\x03\x04":
        return "docx"
    raise AppError(400, "UNSUPPORTED_FILE_TYPE", "Upload a PDF or DOCX file.")


def _finish(units: list, kind: str, ocr_pages: list | None = None, scanned_hint: bool = False) -> dict:
    s = get_settings()
    text = "\n\n".join(u["text"] for u in units)
    if len(text.strip()) < s.min_text_chars:
        hint = ""
        if scanned_hint:
            hint = (" This looks like a scanned PDF. Install Tesseract OCR to read it (see docs/INTEGRATIONS.md), or upload a text-based file." if not ocr_available()
                    else " OCR could not find readable text on these pages.")
        raise AppError(400, "EMPTY_DOCUMENT", "No readable text was found in this file." + hint)
    return {"kind": kind, "text": text, "units": units, "ocr_pages": ocr_pages or []}


def extract(filename: str, data: bytes) -> dict:
    s = get_settings()
    if len(data) > s.max_upload_bytes:
        raise AppError(413, "FILE_TOO_LARGE", f"File is larger than {s.max_upload_mb} MB.")
    kind = _kind(filename, data)
    if kind == "pdf":
        units, ocr_pages, blank = _pdf(data)
        return _finish(units, kind, ocr_pages, scanned_hint=blank > 0)
    return _finish(_docx(data), kind)


def from_text(text: str) -> dict:
    """Plain text (for example a Google Doc export) -> units, one per non-empty line."""
    units = [{"ref": {"paragraph": i + 1}, "text": t.strip()} for i, t in enumerate(text.splitlines()) if t.strip()]
    return _finish(units, "text")


def ocr_available() -> bool:
    if not get_settings().ocr_enabled:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:  # not installed, or the Tesseract program is missing
        return False


def _ocr_page(page) -> str:
    import pytesseract
    from PIL import Image
    pix = page.get_pixmap(dpi=200)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    return pytesseract.image_to_string(img, lang=get_settings().ocr_lang)


def _pdf(data: bytes):
    """Returns (units, pages_read_with_ocr, count_of_blank_pages_not_read)."""
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise AppError(500, "EXTRACTOR_UNAVAILABLE", "PDF support is not installed on the server.") from e
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        pages = [(i + 1, p) for i, p in enumerate(doc)]
    except Exception as e:
        raise AppError(400, "CORRUPT_FILE", "This PDF could not be read.") from e
    units, ocr_pages, blank, can_ocr = [], [], 0, ocr_available()
    for number, page in pages:
        text = (page.get_text() or "").strip()
        if len(text) >= 20:
            units.append({"ref": {"page": number}, "text": text})
        elif can_ocr and len(ocr_pages) < MAX_OCR_PAGES:
            try:
                text = (_ocr_page(page) or "").strip()
            except Exception:  # OCR failure on one page should not fail the whole file
                text = ""
            if text:
                units.append({"ref": {"page": number, "ocr": True}, "text": text})
                ocr_pages.append(number)
            else:
                blank += 1
        else:
            blank += 1
    return units, ocr_pages, blank


def _docx(data: bytes):
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise AppError(400, "CORRUPT_FILE", "This DOCX could not be read.") from e
    units, n = [], 0
    for p in d.paragraphs:
        n += 1
        if p.text.strip():
            units.append({"ref": {"paragraph": n}, "text": p.text.strip()})
    for t in d.tables:
        for row in t.rows:
            line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if line:
                n += 1
                units.append({"ref": {"paragraph": n}, "text": line})
    return units
